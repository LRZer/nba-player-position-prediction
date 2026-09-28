from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from generate_academic_deep_learning_ppt import (
    BG,
    BLUE,
    BLUE_L,
    GREY,
    HIGH,
    INK,
    LINE,
    MID,
    NAVY,
    ORANGE,
    ORANGE_L,
    PANEL,
    RED,
    RED_L,
    RESULTS,
    TEAL,
    TEAL_L,
    WIDE,
    arrow,
    bullets,
    fill_shape,
    image_box,
    metric_card,
    new_slide,
    num,
    panel,
    percent,
    process_box,
    read_json,
    section_label,
    table,
    text_box,
)
from generate_deep_learning_design_ppt import (
    ABLATION_CN,
    COMPARISON_CN,
    save_ablation_cn,
    save_supervised_comparison_cn,
)


OUT_PATH = RESULTS / "NBA_position_deep_learning_teaching_report.pptx"


def callout(slide, title: str, body: list[str], x, y, w, h, accent=TEAL, fill=PANEL, size=11.2):
    panel(slide, x, y, w, h, fill=fill, line=LINE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.055), Inches(h))
    fill_shape(bar, accent)
    text_box(slide, title, x + 0.18, y + 0.14, w - 0.35, 0.26, size=12.2, bold=True, color=NAVY)
    bullets(slide, body, x + 0.20, y + 0.50, w - 0.38, h - 0.60, size=size)


def small_note(slide, text: str, x, y, w, h, fill=TEAL_L, color=NAVY):
    panel(slide, x, y, w, h, fill=fill, line=fill)
    text_box(slide, text, x + 0.12, y + 0.12, w - 0.24, h - 0.20, size=11.2, bold=True, color=color)


def cover(prs: Presentation, dnn: dict, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(WIDE), Inches(0.18))
    fill_shape(band, NAVY)
    text_box(slide, "NBA球员位置预测实验汇报", 0.70, 0.58, 8.70, 0.58, size=29, bold=True, color=NAVY)
    text_box(slide, "重点：深度学习模型设计、模块解释与实验对比", 0.73, 1.18, 9.50, 0.34, size=15, color=GREY)
    text_box(slide, "分类目标：C / PF / PG / SF / SG", 0.75, 1.58, 4.65, 0.24, size=11, bold=True, color=TEAL)
    metric_card(slide, 0.78, 2.25, 2.18, 0.95, "Accuracy", percent(dnn["accuracy"]), "最终测试集", accent=TEAL)
    metric_card(slide, 3.23, 2.25, 2.18, 0.95, "Macro F1", percent(dnn["macro_f1"]), "五类平均", accent=BLUE)
    metric_card(slide, 5.68, 2.25, 2.18, 0.95, "输入特征", "45维", "核心技术画像", accent=ORANGE)
    metric_card(slide, 8.13, 2.25, 2.18, 0.95, "模型框架", "DNN融合", "多模块组合", accent=RED)
    image_box(slide, COMPARISON_CN, 0.78, 3.75, 5.55, 2.60, title="五个实验监督分类对比")
    callout(slide, "本PPT讲解目标", [
        "先说明实验任务和前四个基础模型。",
        "重点解释深度学习模型每个模块的作用。",
        "用测试结果和消融实验说明模型为什么有效。",
    ], 6.75, 3.75, 4.95, 2.55, accent=TEAL, size=11.8)
    text_box(slide, f"01/{total:02d}", 11.85, 7.08, 0.9, 0.22, size=9, color=GREY, align=PP_ALIGN.RIGHT)


def main() -> None:
    bayes = read_json(RESULTS / "experiment1_bayes" / "metrics.json")
    kmeans = read_json(RESULTS / "experiment2_kmeans" / "metrics.json")
    id3 = read_json(RESULTS / "experiment3_id3" / "metrics.json")
    c45 = read_json(RESULTS / "experiment4_c45" / "metrics.json")
    dnn = read_json(RESULTS / "experiment5_dnn" / "metrics.json")
    class_metrics = pd.read_csv(RESULTS / "experiment5_dnn" / "class_metrics.csv")

    save_supervised_comparison_cn(bayes, id3, c45, dnn)
    save_ablation_cn()

    class_rows = []
    for pos in ["C", "PF", "PG", "SF", "SG"]:
        sub = class_metrics[class_metrics["Position"] == pos].set_index("Metric")["Score"]
        class_rows.append([pos, num(float(sub["precision"])), num(float(sub["recall"])), num(float(sub["f1-score"]))])

    prs = Presentation()
    prs.slide_width = Inches(WIDE)
    prs.slide_height = Inches(HIGH)
    total = 22
    page = 1
    cover(prs, dnn, total)
    page += 1

    slide = new_slide(prs, "00 汇报路线", "本次汇报的讲解顺序", "按“任务—对比—模型—结果—消融”的顺序展开，便于理解模型设计", page, total)
    steps = [
        ("1", "任务定义", "说明要预测什么，以及数据来自哪里。"),
        ("2", "基础模型", "用贝叶斯、K-Means、ID3、C4.5建立对照。"),
        ("3", "深度模型", "逐个解释特征、DNN、对比学习、集成和融合。"),
        ("4", "实验结果", "看总体指标、类别表现和混淆矩阵。"),
        ("5", "消融实验", "说明每个模块对最终结果的贡献。"),
    ]
    for i, (idx, title, body) in enumerate(steps):
        y = 1.55 + i * 0.86
        process_box(slide, idx, 0.90, y, 0.45, 0.45, fill=TEAL_L, size=13)
        text_box(slide, title, 1.55, y + 0.03, 1.25, 0.24, size=13.5, bold=True, color=NAVY)
        text_box(slide, body, 3.00, y + 0.03, 8.20, 0.24, size=12.0, color=INK)
    callout(slide, "汇报时的主线", [
        "不是直接说“深度学习更强”，而是说明它解决了什么问题。",
        "前四个模型作为基线，帮助老师看到深度模型的提升幅度。",
        "深度模型部分按模块讲，避免一次性堆出复杂结构。",
    ], 0.95, 6.02, 11.10, 0.95, accent=BLUE, size=10.8)
    page += 1

    slide = new_slide(prs, "01 任务定义", "本实验要解决什么问题", "给定一个球员某赛季的技术统计，预测他更接近五个标准位置中的哪一类", page, total)
    process_box(slide, "输入\n赛季技术统计", 0.90, 2.15, 1.55, 0.80, fill=BLUE_L)
    arrow(slide, 2.55, 2.55, 3.25, 2.55)
    process_box(slide, "模型\n学习统计规律", 3.35, 2.15, 1.65, 0.80, fill=TEAL_L)
    arrow(slide, 5.10, 2.55, 5.80, 2.55)
    process_box(slide, "输出\n五个位置概率", 5.90, 2.15, 1.75, 0.80, fill=BLUE_L)
    arrow(slide, 7.75, 2.55, 8.45, 2.55)
    process_box(slide, "最终类别\n概率最大的位置", 8.55, 2.15, 2.05, 0.80, fill=TEAL_L)
    callout(slide, "五类标签", [
        "C：中锋，通常篮板和封盖更突出。",
        "PF：大前锋，兼具内线和部分锋线特征。",
        "PG：控球后卫，助攻、控球和组织特征更明显。",
        "SF / SG：锋线和得分后卫，边界更连续，容易混淆。",
    ], 0.95, 4.10, 5.25, 1.72, accent=TEAL, size=11.2)
    callout(slide, "任务难点", [
        "NBA位置不是完全离散的规则，很多球员具有摇摆属性。",
        "同一个位置在不同时期的打法也不同，例如现代内线会投更多三分。",
        "因此只靠单条规则很难准确区分所有样本。",
    ], 6.65, 4.10, 5.25, 1.72, accent=ORANGE, size=11.2)
    page += 1

    slide = new_slide(prs, "01 数据处理", "数据流程与合规边界", "模型只使用技术统计，不使用会泄露身份或历史信息的字段", page, total)
    process_box(slide, "原始CSV\nNBA赛季统计", 0.82, 1.82, 1.70, 0.70, fill=BLUE_L)
    arrow(slide, 2.62, 2.17, 3.05, 2.17)
    process_box(slide, "保留标准位置\nC/PF/PG/SF/SG", 3.15, 1.82, 1.95, 0.70, fill=TEAL_L)
    arrow(slide, 5.20, 2.17, 5.63, 2.17)
    process_box(slide, "缺失值处理\n数值统一填充", 5.73, 1.82, 1.85, 0.70, fill=BLUE_L)
    arrow(slide, 7.68, 2.17, 8.10, 2.17)
    process_box(slide, "训练/测试划分\n固定随机种子", 8.20, 1.82, 1.95, 0.70, fill=TEAL_L)
    arrow(slide, 10.25, 2.17, 10.68, 2.17)
    process_box(slide, "统一评估\n输出图表报告", 10.78, 1.82, 1.70, 0.70, fill=BLUE_L)
    callout(slide, "允许使用的信息", [
        "原始技术统计：得分、篮板、助攻、抢断、盖帽、命中率等。",
        "由技术统计派生出的指标：每36分钟数据、比例指标、相邻位置分离指标等。",
        "年份Year可用于刻画不同时代的打法差异。",
    ], 0.95, 3.55, 5.35, 1.62, accent=TEAL, size=11.2)
    callout(slide, "不使用的信息", [
        "Player：球员姓名属于身份信息。",
        "Tm：球队信息可能间接引入队伍风格或角色信息。",
        "历史位置、球员姓名相关统计、球队相关统计等都不加入最终模型。",
    ], 6.65, 3.55, 5.35, 1.62, accent=RED, size=11.2)
    small_note(slide, "这样处理的目的：保证模型确实是根据技术表现判断位置，而不是记住某个球员或球队。", 0.95, 5.85, 11.00, 0.46, fill=ORANGE_L, color=NAVY)
    page += 1

    slide = new_slide(prs, "01 评价指标", "如何判断模型效果", "本实验主要看Accuracy和Macro F1，两者分别反映总体正确率和各类别均衡性", page, total)
    metric_card(slide, 1.10, 1.72, 2.40, 1.08, "Accuracy", "总体正确率", "预测正确样本 / 全部样本", accent=BLUE)
    metric_card(slide, 4.10, 1.72, 2.40, 1.08, "Precision", "预测可靠性", "预测为某类时有多少是真的", accent=TEAL)
    metric_card(slide, 7.10, 1.72, 2.40, 1.08, "Recall", "召回能力", "真实某类中有多少被找出", accent=ORANGE)
    metric_card(slide, 10.10, 1.72, 2.40, 1.08, "F1", "综合指标", "Precision与Recall的平衡", accent=RED)
    callout(slide, "为什么强调Macro F1", [
        "五个位置的样本量和区分难度不完全相同。",
        "Accuracy高可能只是某些大类预测得好。",
        "Macro F1会先分别计算每个类别的F1，再取平均，更能反映五类是否均衡。",
    ], 1.05, 3.55, 5.40, 1.70, accent=BLUE, size=11.5)
    callout(slide, "汇报时可以这样说", [
        "Accuracy回答“总体猜对多少”。",
        "Macro F1回答“五个位置是不是都预测得比较均衡”。",
        "本实验最终模型Accuracy和Macro F1都在0.726左右，说明不是只偏向某一类。",
    ], 6.90, 3.55, 5.10, 1.70, accent=TEAL, size=11.5)
    page += 1

    slide = new_slide(prs, "02 基础模型", "前四个实验为什么作为对照", "基础模型提供可解释基线，但表达能力有限", page, total)
    rows = [
        ["贝叶斯", "假设特征相互独立", "实现简单，但NBA统计特征相关性强"],
        ["K-Means", "按距离把样本分成5簇", "不使用真实标签，适合探索，不适合高精度分类"],
        ["ID3", "按信息增益建立单棵树", "规则清晰，但连续数值离散化会损失细节"],
        ["C4.5", "用信息增益率改进ID3", "仍是单树模型，复杂边界表达不足"],
    ]
    table(slide, rows, ["模型", "核心思路", "主要限制"], 0.90, 1.52, 6.85, 3.00, widths=[1.10, 2.30, 3.45], font_size=9.9)
    callout(slide, "需要说明的重点", [
        "前四个模型不是失败，而是作为基础算法版本进行对照。",
        "它们能说明：只靠简单假设、距离或单棵树，难以充分表达NBA位置边界。",
        "深度学习模型的优势要放在这个对照背景下理解。",
    ], 8.15, 1.72, 3.60, 2.60, accent=ORANGE, size=11.0)
    image_box(slide, RESULTS / "experiment2_kmeans" / "cluster_position_heatmap.png", 8.15, 4.78, 3.60, 1.62, title="K-Means簇与位置混合")
    page += 1

    slide = new_slide(prs, "02 指标对比", "五个实验的整体结果", "深度学习模型在监督分类指标上明显领先", page, total)
    image_box(slide, COMPARISON_CN, 0.85, 1.45, 6.35, 3.72, title="监督分类模型指标对比")
    rows = [
        ["贝叶斯", num(bayes["accuracy"]), num(bayes["macro_f1"])],
        ["ID3", num(id3["accuracy"]), num(id3["macro_f1"])],
        ["C4.5", num(c45["accuracy"]), num(c45["macro_f1"])],
        ["深度学习", num(dnn["accuracy"]), num(dnn["macro_f1"])],
    ]
    table(slide, rows, ["模型", "Accuracy", "Macro F1"], 7.62, 1.70, 3.90, 2.05, widths=[1.35, 1.20, 1.35], font_size=10.4, header_fill=BLUE)
    callout(slide, "结果解释", [
        f"深度学习模型Macro F1为{num(dnn['macro_f1'])}。",
        f"相比贝叶斯提升约{dnn['macro_f1'] - bayes['macro_f1']:.3f}。",
        f"相比ID3提升约{dnn['macro_f1'] - id3['macro_f1']:.3f}。",
        "提升主要来自更强的非线性表示和多模块概率融合。",
    ], 7.62, 4.25, 3.90, 1.55, accent=TEAL, size=10.8)
    page += 1

    slide = new_slide(prs, "03 深度学习动机", "为什么本任务适合深度学习", "NBA位置边界由多个技术指标共同决定，不是单个字段能直接划分", page, total)
    callout(slide, "简单规则的困难", [
        "只看篮板：中锋和大前锋都可能很高。",
        "只看助攻：控卫通常高，但部分前锋也会承担组织。",
        "只看三分：得分后卫、前锋、现代内线都可能有三分。",
    ], 0.95, 1.55, 5.05, 2.00, accent=ORANGE, size=11.5)
    callout(slide, "深度学习的作用", [
        "DNN可以同时利用多个字段，学习“组合关系”。",
        "例如得分、助攻、三分、篮板、防守一起看，位置特征会更清楚。",
        "模型输出的是五个位置的概率，能表达样本本身的模糊性。",
    ], 6.55, 1.55, 5.05, 2.00, accent=TEAL, size=11.5)
    process_box(slide, "得分", 1.35, 4.35, 0.95, 0.46, fill=BLUE_L)
    process_box(slide, "助攻", 2.65, 4.35, 0.95, 0.46, fill=BLUE_L)
    process_box(slide, "篮板", 3.95, 4.35, 0.95, 0.46, fill=BLUE_L)
    process_box(slide, "三分", 5.25, 4.35, 0.95, 0.46, fill=BLUE_L)
    process_box(slide, "防守", 6.55, 4.35, 0.95, 0.46, fill=BLUE_L)
    arrow(slide, 7.70, 4.58, 8.45, 4.58)
    process_box(slide, "DNN学习组合关系", 8.55, 4.25, 1.80, 0.64, fill=TEAL_L)
    arrow(slide, 10.45, 4.58, 11.05, 4.58)
    process_box(slide, "位置概率", 11.15, 4.25, 1.10, 0.64, fill=ORANGE_L, color=ORANGE)
    small_note(slide, "汇报表述：深度学习不是替代数据处理，而是在已有技术统计基础上学习更复杂的组合边界。", 1.00, 5.90, 10.95, 0.46, fill=TEAL_L, color=NAVY)
    page += 1

    slide = new_slide(prs, "04 模型总览", "最终深度学习模型的整体结构", "模型由四类核心模块组成：特征、DNN主干、年代专家、概率融合", page, total)
    process_box(slide, "45维核心特征\n技术统计画像", 0.88, 2.90, 1.65, 0.76, fill=TEAL_L)
    arrow(slide, 2.65, 3.28, 3.10, 3.28)
    process_box(slide, "标准化\n统一量纲", 3.20, 2.90, 1.35, 0.76, fill=BLUE_L)
    arrow(slide, 4.67, 3.28, 5.15, 2.42)
    arrow(slide, 4.67, 3.28, 5.15, 3.28)
    arrow(slide, 4.67, 3.28, 5.15, 4.12)
    process_box(slide, "全局DNN集成\n学习通用规律", 5.25, 2.02, 2.05, 0.76, fill=BLUE_L)
    process_box(slide, "年代专家DNN\n学习时期差异", 5.25, 2.90, 2.05, 0.76, fill=TEAL_L)
    process_box(slide, "ExtraTrees\n辅助概率", 5.25, 3.78, 2.05, 0.76, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 7.42, 2.42, 8.10, 3.08)
    arrow(slide, 7.42, 3.28, 8.10, 3.28)
    arrow(slide, 7.42, 4.15, 8.10, 3.50)
    process_box(slide, "概率融合\n得到最终概率", 8.20, 2.92, 1.85, 0.76, fill=TEAL_L)
    arrow(slide, 10.17, 3.30, 10.70, 3.30)
    process_box(slide, "最终预测\nC/PF/PG/SF/SG", 10.80, 2.92, 1.72, 0.76, fill=BLUE_L)
    callout(slide, "结构解释", [
        "全局DNN学习所有年代、所有样本中的总体规律。",
        "年代专家DNN补充不同时期打法的差异。",
        "ExtraTrees提供不同于神经网络的概率参考。",
        "最后把概率融合，而不是简单投票。",
    ], 1.00, 5.25, 11.00, 1.05, accent=TEAL, size=11.2)
    page += 1

    slide = new_slide(prs, "05 模块一", "输入特征：45维核心技术画像", "特征用于把球员的赛季表现转化为模型可以读取的数字画像", page, total)
    rows = [
        ["基础信息", "Year, Age, G, MP", "描述赛季背景和出场规模"],
        ["效率指标", "FG%, 3P%, eFG%, FT%", "描述投篮效率"],
        ["产量指标", "PPG, RPG, APG, 每36分钟数据", "描述单位时间贡献"],
        ["组织防守", "AST_TOV, STL_36, BLK_36", "描述组织和防守能力"],
        ["位置分离", "Center_PF, SG_PG, SF_PF", "强化相邻位置差异"],
    ]
    table(slide, rows, ["特征类别", "例子", "作用"], 0.88, 1.45, 6.45, 3.05, widths=[1.25, 2.55, 2.65], font_size=9.3, header_fill=TEAL)
    callout(slide, "为什么不用全部字段", [
        "特征过多可能引入噪声，使模型学习到不稳定的细节。",
        "45维版本保留了主要技术画像，同时减少输入维度。",
        "最终效果基本保持，说明筛选没有丢掉关键判别信息。",
    ], 7.82, 1.65, 3.85, 1.90, accent=ORANGE, size=11.0)
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "feature_count_vs_macro_f1.png", 7.82, 4.05, 3.85, 1.95, title="特征数与效果")
    page += 1

    slide = new_slide(prs, "05 模块一", "特征筛选结果怎么解释", "45维正式版在减少输入的同时保持了接近最优的指标", page, total)
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "compact_feature_search_macro_f1.png", 0.85, 1.42, 5.25, 3.50, title="不同紧凑特征集对比")
    table(slide, [
        ["76维全量", "76", "0.7257"],
        ["54维增强", "54", "0.7227"],
        ["45维正式版", "45", "0.7258"],
        ["35维重要性", "35", "0.7158"],
        ["32维画像版", "32", "0.7057"],
    ], ["特征集", "维度", "Macro F1"], 6.70, 1.62, 4.50, 2.30, widths=[1.80, 0.90, 1.35], font_size=10.0, header_fill=BLUE)
    callout(slide, "汇报解释", [
        "45维不是简单取最多特征，而是在效果和简洁性之间折中。",
        "与76维全量相比，45维输入减少约40.8%。",
        "Macro F1仍为0.7258，说明主要位置信息已经被保留下来。",
    ], 6.70, 4.35, 4.50, 1.32, accent=TEAL, size=11.2)
    page += 1

    slide = new_slide(prs, "06 模块二", "Residual DNN主干：模型如何学习", "DNN主干负责把45维特征转化为五个位置的概率", page, total)
    process_box(slide, "输入\n45维特征", 0.95, 2.05, 1.25, 0.70, fill=BLUE_L)
    arrow(slide, 2.30, 2.40, 2.78, 2.40)
    process_box(slide, "线性层\n提取初步表示", 2.88, 1.98, 1.45, 0.84, fill=TEAL_L)
    arrow(slide, 4.43, 2.40, 4.91, 2.40)
    process_box(slide, "残差块\n继续修正特征", 5.00, 1.98, 1.45, 0.84, fill=BLUE_L)
    arrow(slide, 6.55, 2.40, 7.03, 2.40)
    process_box(slide, "分类层\n输出五类概率", 7.12, 1.98, 1.55, 0.84, fill=TEAL_L)
    arrow(slide, 8.78, 2.40, 9.26, 2.40)
    process_box(slide, "预测\n概率最大类别", 9.35, 2.05, 1.55, 0.70, fill=ORANGE_L, color=ORANGE)
    callout(slide, "通俗解释", [
        "线性层可以看作先把原始特征重新组合。",
        "残差块是在已有组合基础上继续做修正。",
        "分类层把修正后的表示转换成C、PF、PG、SF、SG五个概率。",
    ], 1.05, 4.10, 5.30, 1.48, accent=TEAL, size=11.2)
    callout(slide, "为什么适合本任务", [
        "位置判断通常不是单一字段决定，而是多个指标共同决定。",
        "DNN可以学习“得分+助攻+篮板+效率”等组合关系。",
        "概率输出也能表达摇摆位置的不确定性。",
    ], 6.75, 4.10, 5.05, 1.48, accent=BLUE, size=11.2)
    page += 1

    slide = new_slide(prs, "06 模块二", "残差连接：为什么叫Residual", "残差结构的核心是保留原始信息，再学习需要补充的修正量", page, total)
    process_box(slide, "输入 x", 1.05, 2.30, 1.00, 0.55, fill=BLUE_L)
    arrow(slide, 2.12, 2.58, 2.85, 2.58)
    process_box(slide, "普通网络层\n学习 F(x)", 2.95, 2.13, 1.65, 0.88, fill=TEAL_L)
    arrow(slide, 4.70, 2.58, 5.45, 2.58)
    process_box(slide, "相加\nx + F(x)", 5.55, 2.13, 1.35, 0.88, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 6.98, 2.58, 7.72, 2.58)
    process_box(slide, "输出\n更新后的表示", 7.82, 2.30, 1.55, 0.55, fill=BLUE_L)
    callout(slide, "可以这样理解", [
        "模型不是完全重新写一个表示，而是在原表示上做修正。",
        "如果某一层暂时学不到有用信息，原来的x仍然可以保留下来。",
        "这样网络更容易训练，训练过程更稳定。",
    ], 1.05, 4.05, 5.00, 1.45, accent=TEAL, size=11.2)
    callout(slide, "在实验中的设置", [
        "Residual DNN使用宽度为192的隐藏表示。",
        "残差块数量为2个，并配合BatchNorm、ReLU和Dropout。",
        "Dropout用于减少模型过拟合。",
    ], 6.65, 4.05, 5.00, 1.45, accent=BLUE, size=11.2)
    page += 1

    slide = new_slide(prs, "07 模块三", "监督对比学习：让隐藏表示更清楚", "该模块帮助模型在内部空间中把相同位置样本放近、不同位置样本拉远", page, total)
    panel(slide, 0.95, 1.60, 4.10, 3.00, fill=PANEL, line=LINE)
    text_box(slide, "隐藏空间示意", 1.20, 1.85, 1.50, 0.24, size=12.0, bold=True, color=NAVY)
    for x, y, color in [
        (1.65, 2.45, TEAL), (1.95, 2.20, TEAL), (2.20, 2.55, TEAL),
        (3.35, 3.15, BLUE), (3.65, 3.40, BLUE), (3.15, 3.50, BLUE),
        (2.85, 2.68, ORANGE), (3.05, 2.35, ORANGE),
    ]:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.16), Inches(0.16))
        fill_shape(dot, color)
    text_box(slide, "同色点代表同一位置", 1.45, 4.10, 2.60, 0.20, size=9.2, color=GREY, align=PP_ALIGN.CENTER)
    callout(slide, "训练目标", [
        "交叉熵损失：要求最终预测类别正确。",
        "监督对比损失：要求隐藏表示按真实类别更好地分开。",
        "总损失 = 交叉熵 + 0.1 × 监督对比损失。",
    ], 5.65, 1.62, 5.35, 1.70, accent=ORANGE, size=11.4)
    callout(slide, "为什么有用", [
        "C/PF、PF/SF、PG/SG这类相邻位置容易混在一起。",
        "监督对比学习相当于给模型一个额外约束：同类更接近，异类更分开。",
        "它不改变标签，也不加入新的身份信息。",
    ], 5.65, 3.82, 5.35, 1.70, accent=TEAL, size=11.4)
    page += 1

    slide = new_slide(prs, "08 模块四", "全局DNN集成：多个模型共同判断", "单个神经网络会受随机初始化影响，集成可以让预测概率更稳定", page, total)
    for i, x in enumerate([0.95, 2.45, 3.95, 5.45, 6.95, 8.45]):
        fill = BLUE_L if i < 3 else TEAL_L
        label = "Residual\nDNN" if i < 3 else "Contrastive\nDNN"
        process_box(slide, f"{label}\n成员{i + 1}", x, 2.00, 1.05, 0.85, fill=fill, size=9.2)
        arrow(slide, x + 0.52, 2.93, 5.70, 4.02, color=MID, width=0.8)
    process_box(slide, "平均/加权平均\n得到全局DNN概率", 4.70, 4.10, 2.05, 0.78, fill=ORANGE_L, color=ORANGE)
    callout(slide, "通俗解释", [
        "不是只让一个模型做决定，而是让多个DNN分别给出概率。",
        "多个模型的概率再做平均，可以减少单个模型的偶然误差。",
        "这类似多个评委独立打分后取综合结果。",
    ], 0.95, 5.40, 5.15, 1.10, accent=TEAL, size=11.2)
    callout(slide, "实验依据", [
        "完整模型使用多个DNN成员。",
        "消融实验中退化为单个Residual DNN后，Macro F1下降约0.0340。",
        "说明集成是本模型的重要组成部分。",
    ], 6.55, 5.40, 5.15, 1.10, accent=BLUE, size=11.2)
    page += 1

    slide = new_slide(prs, "09 模块五", "年代专家DNN：处理时代差异", "NBA不同年代的位置职责不同，年代专家用于学习更局部的规律", page, total)
    for i, (era, x) in enumerate([("1980s", 1.00), ("1990s", 3.00), ("2000s", 5.00), ("2010s+", 7.00)]):
        process_box(slide, f"{era}\n专家DNN", x, 2.05, 1.35, 0.82, fill=TEAL_L if i % 2 else BLUE_L)
        arrow(slide, x + 0.68, 2.95, 5.70, 4.05, color=MID, width=0.8)
    process_box(slide, "对应年代的专家概率\nP_era", 4.78, 4.12, 2.00, 0.78, fill=ORANGE_L, color=ORANGE)
    callout(slide, "为什么要分年代", [
        "同样是内线球员，早期更偏篮下，现代可能有更多三分。",
        "后卫和锋线的组织、投篮职责也随时代变化。",
        "Year不是身份信息，而是描述赛季背景的技术环境。",
    ], 0.95, 5.28, 5.20, 1.12, accent=TEAL, size=11.1)
    callout(slide, "实验依据", [
        "年代专家按Year划分为多个时期。",
        "去掉年代专家后，Macro F1下降约0.0176。",
        "说明时代差异确实为位置判断提供了信息。",
    ], 6.55, 5.28, 5.20, 1.12, accent=BLUE, size=11.1)
    page += 1

    slide = new_slide(prs, "10 模块六", "ExtraTrees概率融合：补充另一种判断方式", "ExtraTrees是树模型，它单独不如DNN，但能提供不同的概率参考", page, total)
    process_box(slide, "DNN概率\nP_dnn", 1.00, 2.10, 1.45, 0.72, fill=BLUE_L)
    process_box(slide, "ExtraTrees概率\nP_tree", 1.00, 3.22, 1.45, 0.72, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 2.58, 2.46, 4.10, 3.00)
    arrow(slide, 2.58, 3.58, 4.10, 3.30)
    process_box(slide, "加权融合\nP_final = 0.55P_dnn + 0.45P_tree", 4.20, 2.85, 3.05, 0.86, fill=TEAL_L)
    arrow(slide, 7.38, 3.28, 8.10, 3.28)
    process_box(slide, "选择概率最大的位置", 8.20, 2.96, 2.00, 0.64, fill=BLUE_L)
    callout(slide, "为什么不是硬投票", [
        "硬投票只看最终类别，丢掉了概率大小。",
        "概率融合可以保留模型对每个类别的不确定性。",
        "对于PF/SF/SG这类模糊位置，概率信息更有价值。",
    ], 1.00, 4.85, 5.10, 1.22, accent=TEAL, size=11.0)
    callout(slide, "实验依据", [
        "ExtraTrees单独效果低于DNN，不作为最终主模型。",
        "去掉ExtraTrees融合后，Macro F1下降约0.0045。",
        "说明它作为辅助概率来源有小幅稳定贡献。",
    ], 6.58, 4.85, 5.10, 1.22, accent=ORANGE, size=11.0)
    page += 1

    slide = new_slide(prs, "11 训练流程", "模型如何训练和验证", "训练流程保证模型配置不是直接在测试集上挑出来的", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "training_curve.png", 0.85, 1.45, 5.25, 3.40, title="训练曲线")
    callout(slide, "训练设置", [
        "框架：PyTorch，设备：CUDA GPU。",
        "优化器：AdamW，用于更新神经网络参数。",
        "Batch size：512，每次用一批样本训练。",
        "Early stopping：验证集长期不提升时停止训练。",
    ], 6.65, 1.62, 4.85, 2.05, accent=BLUE, size=11.0)
    callout(slide, "验证流程", [
        "先从训练集中划分验证集，用来比较配置。",
        "配置确定后，再使用训练数据重训最终模型。",
        "最后只在测试集上统一评估一次。",
    ], 6.65, 4.05, 4.85, 1.34, accent=TEAL, size=11.0)
    page += 1

    slide = new_slide(prs, "12 测试结果", "最终模型的分类效果", "混淆矩阵展示每个真实位置被预测成各类别的数量", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "confusion_matrix.png", 0.80, 1.42, 5.60, 4.10, title="深度学习模型混淆矩阵")
    table(slide, class_rows, ["位置", "Precision", "Recall", "F1"], 6.85, 1.58, 4.58, 2.46, widths=[0.72, 1.22, 1.22, 1.22], font_size=10.0, header_fill=TEAL)
    callout(slide, "读图方式", [
        "对角线表示预测正确的样本数量。",
        "非对角线表示混淆，例如真实SF被预测成SG。",
        "PG识别最稳定，PF/SF/SG之间混淆更多。",
    ], 6.85, 4.45, 4.58, 1.20, accent=BLUE, size=11.0)
    page += 1

    slide = new_slide(prs, "12 测试结果", "错误模式如何解释", "模型的主要错误集中在篮球语义上相邻的位置之间", page, total)
    callout(slide, "主要混淆一：内线边界", [
        "C → PF：162个样本。",
        "PF → C：148个样本。",
        "原因：中锋和大前锋在篮板、防守、内线得分上有重叠。",
    ], 0.95, 1.58, 5.10, 1.55, accent=BLUE, size=11.2)
    callout(slide, "主要混淆二：锋卫摇摆", [
        "SF → SG：118个样本。",
        "SG → SF：118个样本。",
        "原因：锋线和得分后卫都可能承担得分、三分和部分组织任务。",
    ], 6.60, 1.58, 5.10, 1.55, accent=TEAL, size=11.2)
    callout(slide, "主要混淆三：后卫边界", [
        "PG → SG：88个样本。",
        "SG → PG：86个样本。",
        "原因：得分型控卫和组织型分卫在统计上较接近。",
    ], 0.95, 3.82, 5.10, 1.55, accent=ORANGE, size=11.2)
    callout(slide, "结论", [
        "错误不是随机分散，而是集中在相邻位置。",
        "这说明模型学到的边界与篮球位置语义基本一致。",
        "后续提升空间主要在摇摆位置和多职责球员。",
    ], 6.60, 3.82, 5.10, 1.55, accent=RED, size=11.2)
    page += 1

    slide = new_slide(prs, "13 消融实验", "每个模块到底贡献多少", "消融实验通过“移除一个模块再评估”来验证模块价值", page, total)
    image_box(slide, ABLATION_CN, 0.85, 1.42, 6.70, 4.20, title="中文消融实验图")
    callout(slide, "怎么读这张图", [
        "横轴是移除模块后Macro F1下降的幅度。",
        "条越长，说明该模块越重要。",
        "DNN主干、DNN集成、工程特征、年代专家贡献最明显。",
    ], 8.05, 1.70, 3.55, 1.65, accent=TEAL, size=11.0)
    callout(slide, "实验结论", [
        "完整模型不是靠单个技巧提升。",
        "主要提升来自DNN表示能力和集成稳定性。",
        "工程特征、年代专家和概率融合提供进一步补充。",
    ], 8.05, 3.85, 3.55, 1.65, accent=BLUE, size=11.0)
    page += 1

    slide = new_slide(prs, "14 总结", "最终结论与汇报收束", "深度学习模型在合规约束下取得了五个实验中的最好分类效果", page, total)
    metric_card(slide, 0.90, 1.55, 2.15, 0.90, "Accuracy", percent(dnn["accuracy"]), "测试集", accent=TEAL)
    metric_card(slide, 3.35, 1.55, 2.15, 0.90, "Macro F1", percent(dnn["macro_f1"]), "五类平均", accent=BLUE)
    metric_card(slide, 5.80, 1.55, 2.15, 0.90, "特征数", "45维", "正式版本", accent=ORANGE)
    metric_card(slide, 8.25, 1.55, 2.15, 0.90, "合规性", "满足", "无身份字段", accent=RED)
    callout(slide, "最终模型组成", [
        "45维核心技术特征提供输入画像。",
        "Residual DNN学习复杂非线性位置边界。",
        "监督对比学习帮助相邻类别在隐藏空间中分开。",
        "全局集成、年代专家和ExtraTrees概率融合提升稳定性。",
    ], 0.95, 3.05, 5.30, 2.05, accent=TEAL, size=11.2)
    callout(slide, "答辩时的收束表述", [
        "前四个基础模型提供了对照，说明简单假设和单树规则存在表达限制。",
        "实验五通过特征工程和深度学习模块组合，更好处理NBA位置的连续边界。",
        "消融实验说明各模块都有可解释的贡献，最终模型效果最优。",
    ], 6.65, 3.05, 5.30, 2.05, accent=BLUE, size=11.2)
    small_note(slide, "一句话总结：本模型的优势来自“合规技术特征 + 非线性表示 + 多模型概率融合”，而不是使用身份信息。", 0.95, 5.95, 11.00, 0.46, fill=ORANGE_L, color=NAVY)

    prs.save(OUT_PATH)
    print(f"Saved teaching PPT: {OUT_PATH}")


if __name__ == "__main__":
    main()
