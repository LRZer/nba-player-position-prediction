from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.font_manager import FontProperties
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from generate_academic_deep_learning_ppt import (
    BG,
    BLUE,
    BLUE_L,
    GREY,
    HIGH,
    INK,
    LINE,
    MID,
    NAVY,
    ORANGE,
    ORANGE_L,
    PANEL,
    RED,
    RED_L,
    RESULTS,
    ROOT,
    TEAL,
    TEAL_L,
    WIDE,
    arrow,
    bullets,
    fill_shape,
    image_box,
    metric_card,
    new_slide,
    num,
    panel,
    percent,
    process_box,
    read_json,
    section_label,
    table,
    text_box,
)


OUT_PATH = RESULTS / "NBA_position_deep_learning_design_report.pptx"
COMPARISON_CN = RESULTS / "supervised_model_comparison_cn.png"
ABLATION_CN = RESULTS / "experiment5_ablation" / "ablation_delta_macro_f1_cn.png"


def font_prop() -> FontProperties | None:
    for path in [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]:
        if path.exists():
            return FontProperties(fname=str(path))
    return None


def save_supervised_comparison_cn(bayes: dict, id3: dict, c45: dict, dnn: dict) -> None:
    fp = font_prop()
    labels = ["贝叶斯", "ID3", "C4.5", "深度学习"]
    acc = [bayes["accuracy"], id3["accuracy"], c45["accuracy"], dnn["accuracy"]]
    f1 = [bayes["macro_f1"], id3["macro_f1"], c45["macro_f1"], dnn["macro_f1"]]

    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(8.6, 4.6), dpi=180)
    x = range(len(labels))
    width = 0.34
    bars1 = ax.bar([i - width / 2 for i in x], acc, width, label="Accuracy", color="#2463AB")
    bars2 = ax.bar([i + width / 2 for i in x], f1, width, label="Macro F1", color="#1A907F")
    ax.set_ylim(0, 0.82)
    ax.set_ylabel("指标值", fontproperties=fp, fontsize=11)
    ax.set_title("监督分类模型指标对比", fontproperties=fp, fontsize=14, pad=12)
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontproperties=fp, fontsize=11)
    ax.grid(axis="y", alpha=0.25)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(prop=fp, frameon=False, loc="upper left")
    for bars in [bars1, bars2]:
        for bar in bars:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 0.015,
                f"{bar.get_height():.3f}",
                ha="center",
                va="bottom",
                fontsize=9,
                color="#1F2937",
                fontproperties=fp,
            )
    fig.tight_layout()
    fig.savefig(COMPARISON_CN, bbox_inches="tight")
    plt.close(fig)


def save_ablation_cn() -> None:
    fp = font_prop()
    rows = [
        ("只用树模型\n不使用DNN", 0.0393),
        ("单个DNN\n不做集成", 0.0340),
        ("去掉工程特征", 0.0235),
        ("去掉年代专家", 0.0176),
        ("去掉ExtraTrees融合", 0.0045),
        ("去掉对比DNN成员", 0.0045),
        ("去掉对比学习损失", 0.0011),
    ]
    labels = [r[0] for r in rows][::-1]
    values = [r[1] for r in rows][::-1]

    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(8.2, 4.8), dpi=180)
    colors = ["#1A907F" if v >= 0.015 else "#90B4D8" for v in values]
    bars = ax.barh(labels, values, color=colors, height=0.58)
    ax.set_xlabel("Macro F1下降值", fontproperties=fp, fontsize=11)
    ax.set_title("消融实验：移除模块后的性能下降", fontproperties=fp, fontsize=14, pad=12)
    ax.set_xlim(0, 0.043)
    ax.grid(axis="x", alpha=0.25)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    for tick in ax.get_yticklabels():
        tick.set_fontproperties(fp)
        tick.set_fontsize(10.5)
    for bar in bars:
        ax.text(
            bar.get_width() + 0.0008,
            bar.get_y() + bar.get_height() / 2,
            f"{bar.get_width():.4f}",
            va="center",
            fontsize=9.5,
            color="#1F2937",
            fontproperties=fp,
        )
    fig.tight_layout()
    ABLATION_CN.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(ABLATION_CN, bbox_inches="tight")
    plt.close(fig)


def cover_slide(prs: Presentation, dnn: dict, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(WIDE), Inches(0.18))
    fill_shape(band, NAVY)
    text_box(slide, "NBA球员位置预测：深度学习模型设计与实验对比", 0.70, 0.62, 11.3, 0.60, size=27, bold=True, color=NAVY)
    text_box(slide, "基于赛季技术统计的五分类任务 · 重点讲解实验五模型结构", 0.73, 1.20, 10.0, 0.30, size=14.5, color=GREY)
    text_box(slide, "C / PF / PG / SF / SG", 0.75, 1.58, 4.0, 0.24, size=10.5, bold=True, color=TEAL)
    metric_card(slide, 0.78, 2.25, 2.25, 0.95, "最终 Accuracy", percent(dnn["accuracy"]), "测试集", accent=TEAL)
    metric_card(slide, 3.30, 2.25, 2.25, 0.95, "最终 Macro F1", percent(dnn["macro_f1"]), "五类平均", accent=BLUE)
    metric_card(slide, 5.82, 2.25, 2.25, 0.95, "正式输入特征", "45维", "核心技术特征", accent=ORANGE)
    metric_card(slide, 8.34, 2.25, 2.25, 0.95, "合规约束", "无身份字段", "不使用Player/Tm", accent=RED)
    image_box(slide, COMPARISON_CN, 0.78, 3.72, 5.45, 2.62, title="五个实验中的监督分类对比")
    panel(slide, 6.65, 3.72, 5.72, 2.62, fill=PANEL, line=LINE)
    section_label(slide, "本次汇报重点", 6.92, 4.05, 1.38, fill=TEAL_L, color=TEAL)
    bullets(slide, [
        "深度学习模型如何接收45维技术特征",
        "Residual DNN如何学习非线性位置边界",
        "年代专家和概率融合如何补充主模型",
        "用结果图和消融实验证明各模块作用",
    ], 6.95, 4.52, 5.05, 1.05, size=12.8)
    text_box(slide, f"01/{total:02d}", 11.85, 7.08, 0.9, 0.22, size=9, color=GREY, align=PP_ALIGN.RIGHT)


def main() -> None:
    bayes = read_json(RESULTS / "experiment1_bayes" / "metrics.json")
    kmeans = read_json(RESULTS / "experiment2_kmeans" / "metrics.json")
    id3 = read_json(RESULTS / "experiment3_id3" / "metrics.json")
    c45 = read_json(RESULTS / "experiment4_c45" / "metrics.json")
    dnn = read_json(RESULTS / "experiment5_dnn" / "metrics.json")
    class_metrics = pd.read_csv(RESULTS / "experiment5_dnn" / "class_metrics.csv")

    save_supervised_comparison_cn(bayes, id3, c45, dnn)
    save_ablation_cn()

    class_rows = []
    for pos in ["C", "PF", "PG", "SF", "SG"]:
        sub = class_metrics[class_metrics["Position"] == pos].set_index("Metric")["Score"]
        class_rows.append([pos, num(float(sub["precision"])), num(float(sub["recall"])), num(float(sub["f1-score"]))])

    prs = Presentation()
    prs.slide_width = Inches(WIDE)
    prs.slide_height = Inches(HIGH)
    total = 14
    page = 1
    cover_slide(prs, dnn, total)
    page += 1

    slide = new_slide(prs, "01 任务说明", "任务定义与数据流程", "目标是用NBA球员赛季技术统计预测五个标准场上位置", page, total)
    process_box(slide, "原始数据\n赛季技术统计", 0.82, 1.85, 1.85, 0.72, fill=BLUE_L)
    arrow(slide, 2.75, 2.20, 3.12, 2.20)
    process_box(slide, "清洗\n去除身份字段", 3.20, 1.85, 1.85, 0.72, fill=TEAL_L)
    arrow(slide, 5.13, 2.20, 5.50, 2.20)
    process_box(slide, "特征\n传统26维 / DNN45维", 5.58, 1.85, 2.15, 0.72, fill=BLUE_L)
    arrow(slide, 7.82, 2.20, 8.18, 2.20)
    process_box(slide, "训练\n五个实验", 8.26, 1.85, 1.72, 0.72, fill=TEAL_L)
    arrow(slide, 10.07, 2.20, 10.44, 2.20)
    process_box(slide, "评估\nAccuracy / F1", 10.52, 1.85, 1.82, 0.72, fill=BLUE_L)
    metric_card(slide, 1.00, 3.45, 2.25, 0.88, "分类目标", "5类", "C/PF/PG/SF/SG", accent=BLUE)
    metric_card(slide, 3.60, 3.45, 2.25, 0.88, "训练样本", "14,981", "外层训练集", accent=TEAL)
    metric_card(slide, 6.20, 3.45, 2.25, 0.88, "测试样本", "3,746", "统一评估", accent=BLUE)
    metric_card(slide, 8.80, 3.45, 2.25, 0.88, "DNN验证集", "2,997", "训练集内划分", accent=TEAL)
    bullets(slide, [
        "前四个实验保持基础算法版本，作为对照组。",
        "实验五只使用技术统计和由技术统计派生的指标。",
        "不使用 Player、Tm、历史位置等可能造成泄露的信息。",
    ], 1.02, 5.15, 10.9, 1.02, size=13.2)
    page += 1

    slide = new_slide(prs, "02 基线对比", "前四个基础模型的作用", "这些模型用于形成清晰对照，说明深度学习模型的提升来自更强的表示能力", page, total)
    table(slide, [
        ["贝叶斯", "条件独立假设", num(bayes["accuracy"]), num(bayes["macro_f1"])],
        ["K-Means", "距离聚类，不使用标签", "-", f"ARI {num(kmeans['adjusted_rand_index'])}"],
        ["ID3", "信息增益单树", num(id3["accuracy"]), num(id3["macro_f1"])],
        ["C4.5", "信息增益率单树", num(c45["accuracy"]), num(c45["macro_f1"])],
        ["深度学习", "非线性表示 + 概率融合", num(dnn["accuracy"]), num(dnn["macro_f1"])],
    ], ["模型", "核心思想", "Accuracy", "Macro F1"], 0.78, 1.58, 5.35, 2.82, widths=[1.05, 2.05, 1.05, 1.20], font_size=9.7)
    image_box(slide, COMPARISON_CN, 6.65, 1.48, 5.15, 3.05, title="监督分类模型指标对比")
    bullets(slide, [
        "贝叶斯和单树模型更依赖明确假设，难以表达复杂连续边界。",
        "K-Means是无监督聚类，适合探索结构，但不是高精度分类器。",
        "深度学习模型把多个弱边界信号组合起来，整体指标明显更高。",
    ], 0.95, 5.25, 11.0, 1.0, size=12.8)
    page += 1

    slide = new_slide(prs, "03 模型总览", "实验五深度学习模型的整体设计", "先由DNN学习主要分类边界，再用年代专家和树模型概率进行补充", page, total)
    process_box(slide, "45维\n核心技术特征", 0.78, 2.85, 1.62, 0.76, fill=TEAL_L)
    arrow(slide, 2.48, 3.23, 2.92, 3.23)
    process_box(slide, "标准化\n统一量纲", 3.00, 2.85, 1.35, 0.76, fill=BLUE_L)
    arrow(slide, 4.43, 3.23, 4.87, 2.40)
    arrow(slide, 4.43, 3.23, 4.87, 3.23)
    arrow(slide, 4.43, 3.23, 4.87, 4.05)
    process_box(slide, "全局DNN集成\n学习通用边界", 4.98, 2.00, 2.00, 0.76, fill=BLUE_L)
    process_box(slide, "年代专家DNN\n学习年代差异", 4.98, 2.85, 2.00, 0.76, fill=TEAL_L)
    process_box(slide, "ExtraTrees\n补充树模型概率", 4.98, 3.70, 2.00, 0.76, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 7.08, 2.38, 7.75, 3.05)
    arrow(slide, 7.08, 3.23, 7.75, 3.23)
    arrow(slide, 7.08, 4.08, 7.75, 3.42)
    process_box(slide, "概率融合\n保留不确定性", 7.86, 2.88, 1.80, 0.76, fill=TEAL_L)
    arrow(slide, 9.74, 3.23, 10.18, 3.23)
    process_box(slide, "最终预测\n五类位置", 10.28, 2.88, 1.64, 0.76, fill=BLUE_L)
    bullets(slide, [
        "DNN主干负责学习复杂的统计组合关系。",
        "年代专家负责处理不同时期位置打法差异。",
        "ExtraTrees不是替代DNN，而是提供另一种概率参考。",
        "最后使用概率融合，而不是简单投票。",
    ], 0.95, 5.20, 11.15, 1.02, size=12.8)
    page += 1

    slide = new_slide(prs, "04 模块一", "输入特征：45维核心技术画像", "特征不追求越多越好，目标是保留最能区分位置的技术信号", page, total)
    families = [
        ("基础表现", "Year, Age, G, MP"),
        ("进攻效率", "FG%, 3P%, 2P%, eFG%, FT%"),
        ("产量与节奏", "PPG, MPG, PTS_36, FTA_36"),
        ("篮板与防守", "ORB, DRB, TRB, STL, BLK"),
        ("组织与失误", "AST, TOV, APG, AST_TOV"),
        ("位置分离指标", "Center_PF, SF_PF, SG_PG"),
    ]
    for i, (name, desc) in enumerate(families):
        x = 0.85 + (i % 2) * 3.00
        y = 1.62 + (i // 2) * 0.84
        panel(slide, x, y, 2.62, 0.56, fill=TEAL_L if i % 2 == 0 else BLUE_L, line=LINE)
        text_box(slide, name, x + 0.12, y + 0.11, 0.88, 0.20, size=9.2, bold=True, color=NAVY)
        text_box(slide, desc, x + 1.00, y + 0.11, 1.45, 0.20, size=8.6, color=INK)
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "feature_count_vs_macro_f1.png", 7.00, 1.42, 4.70, 3.28, title="特征数量与模型效果")
    bullets(slide, [
        "45维版本从76维候选特征中筛选得到，减少约40.8%的输入。",
        "保留效率、产量、组织、防守、篮板和相邻位置分离指标。",
        "结果基本不下降：Accuracy 0.7261，Macro F1 0.7258。",
    ], 0.95, 5.25, 11.0, 0.96, size=12.8)
    page += 1

    slide = new_slide(prs, "05 模块二", "Residual DNN主干：学习非线性边界", "普通线性规则难以表达位置边界，DNN用于学习多项技术统计的组合关系", page, total)
    process_box(slide, "输入\n45维特征", 0.95, 2.05, 1.20, 0.70, fill=BLUE_L)
    arrow(slide, 2.22, 2.40, 2.70, 2.40)
    process_box(slide, "线性层\n+ BN + ReLU", 2.78, 1.98, 1.45, 0.84, fill=TEAL_L)
    arrow(slide, 4.30, 2.40, 4.78, 2.40)
    process_box(slide, "残差块 ×2\nx + F(x)", 4.86, 1.98, 1.45, 0.84, fill=BLUE_L)
    arrow(slide, 6.38, 2.40, 6.86, 2.40)
    process_box(slide, "分类层\n输出5类概率", 6.94, 1.98, 1.55, 0.84, fill=TEAL_L)
    arrow(slide, 8.56, 2.40, 9.04, 2.40)
    process_box(slide, "预测位置\nC/PF/PG/SF/SG", 9.12, 2.05, 1.70, 0.70, fill=BLUE_L)
    panel(slide, 1.05, 4.02, 4.85, 1.35, fill=PANEL, line=LINE)
    text_box(slide, "残差连接怎么理解", 1.25, 4.22, 2.50, 0.24, size=12.5, bold=True, color=NAVY)
    bullets(slide, [
        "普通层学习 F(x)，残差层学习 x + F(x)。",
        "保留原始信息，再叠加模型学到的修正。",
        "这样训练更稳定，也更适合加深网络。",
    ], 1.25, 4.62, 4.25, 0.72, size=11.5)
    panel(slide, 6.38, 4.02, 4.85, 1.35, fill=PANEL, line=LINE)
    text_box(slide, "在本任务中的作用", 6.58, 4.22, 2.50, 0.24, size=12.5, bold=True, color=NAVY)
    bullets(slide, [
        "组合得分、篮板、助攻、防守等指标。",
        "学习C/PF、PF/SF、PG/SG等相邻边界。",
        "输出每个位置的概率，而不是硬规则。",
    ], 6.58, 4.62, 4.25, 0.72, size=11.5)
    page += 1

    slide = new_slide(prs, "06 模块三", "监督对比学习：让相同位置更接近", "该模块不是单独分类器，而是帮助DNN的隐藏表示更容易区分类别", page, total)
    process_box(slide, "同类样本\n距离拉近", 1.10, 1.80, 1.62, 0.62, fill=TEAL_L)
    process_box(slide, "不同类样本\n距离拉远", 1.10, 3.05, 1.62, 0.62, fill=ORANGE_L, color=ORANGE)
    panel(slide, 3.30, 1.55, 3.20, 2.58, fill=PANEL, line=LINE)
    text_box(slide, "隐空间示意", 3.54, 1.80, 1.20, 0.22, size=11.5, bold=True, color=NAVY)
    for x, y, color in [(4.0, 2.35, TEAL), (4.35, 2.12, TEAL), (4.55, 2.45, TEAL), (5.45, 3.03, BLUE), (5.78, 3.28, BLUE), (5.20, 3.38, BLUE), (4.88, 2.76, ORANGE), (5.03, 2.45, ORANGE)]:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.16), Inches(0.16))
        fill_shape(dot, color)
    text_box(slide, "C/PF/PG/SF/SG", 4.03, 3.78, 1.95, 0.18, size=8.8, color=GREY, align=PP_ALIGN.CENTER)
    panel(slide, 7.05, 1.55, 4.40, 2.58, fill=PANEL, line=LINE)
    text_box(slide, "训练目标", 7.30, 1.82, 1.20, 0.22, size=11.5, bold=True, color=NAVY)
    text_box(slide, "总损失 = 交叉熵损失 + 0.1 × 监督对比损失", 7.30, 2.32, 3.75, 0.28, size=13.5, bold=True, color=NAVY)
    bullets(slide, [
        "交叉熵：要求最终类别预测正确。",
        "监督对比：要求隐藏表示按真实位置形成更清晰的分布。",
        "作用重点：缓解相邻位置样本混在一起的问题。",
    ], 7.30, 2.92, 3.72, 0.88, size=11.4)
    bullets(slide, [
        "解释时可以说：这一步是在训练过程中给模型增加“同类靠近、异类分开”的约束。",
        "它不会改变任务标签，也不引入额外身份信息。",
    ], 1.10, 5.10, 10.45, 0.80, size=12.5)
    page += 1

    slide = new_slide(prs, "07 模块四", "全局DNN集成：降低单模型偶然性", "同一结构在不同随机种子下会有轻微差异，集成用于获得更稳定的概率", page, total)
    for i, x in enumerate([1.10, 2.75, 4.40, 6.05, 7.70, 9.35]):
        fill = BLUE_L if i < 3 else TEAL_L
        label = "Residual\nDNN" if i < 3 else "Contrastive\nDNN"
        process_box(slide, f"{label}\nseed {i + 1}", x, 2.00, 1.18, 0.92, fill=fill, size=9.5)
        arrow(slide, x + 0.58, 2.98, 6.20, 4.02, color=MID, width=0.9)
    process_box(slide, "平均概率\n而不是投票", 5.28, 4.10, 1.85, 0.78, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 7.22, 4.49, 8.10, 4.49)
    process_box(slide, "全局DNN概率\nP_global", 8.20, 4.10, 1.72, 0.78, fill=TEAL_L)
    bullets(slide, [
        "6个成员：3个普通Residual DNN，3个Contrastive Residual DNN。",
        "每个成员输出五个位置的概率，最后取加权平均。",
        "消融结果：退化成单个Residual DNN后，Macro F1下降约0.0340。",
    ], 1.05, 5.55, 10.80, 0.84, size=12.5)
    page += 1

    slide = new_slide(prs, "08 模块五", "年代专家DNN：处理不同时期打法差异", "NBA位置职责随年代变化，分年代训练专家模型可以学习更局部的边界", page, total)
    for i, (era, x) in enumerate([("1980s", 0.95), ("1990s", 3.05), ("2000s", 5.15), ("2010s+", 7.25)]):
        process_box(slide, f"{era}\n专家DNN", x, 2.02, 1.45, 0.82, fill=TEAL_L if i % 2 else BLUE_L)
        arrow(slide, x + 0.72, 2.92, 5.82, 4.02, color=MID, width=0.9)
    process_box(slide, "按Year选择\n对应年代专家", 9.40, 2.02, 1.75, 0.82, fill=ORANGE_L, color=ORANGE)
    process_box(slide, "专家概率\nP_era", 5.08, 4.10, 1.82, 0.78, fill=TEAL_L)
    bullets(slide, [
        "早期内线和现代空间型内线的技术统计边界不同。",
        "后卫、锋线在三分出手和助攻方式上也随年代变化。",
        "年代专家只使用Year分段，不使用球员身份或历史位置。",
        "消融结果：去掉年代专家后，Macro F1下降约0.0176。",
    ], 1.05, 5.28, 10.90, 1.02, size=12.5)
    page += 1

    slide = new_slide(prs, "09 模块六", "ExtraTrees概率融合：补充DNN的错误模式", "树模型单独效果低于DNN，但它的错误模式不同，可以作为概率补充", page, total)
    process_box(slide, "DNN概率\nP_dnn", 1.18, 2.20, 1.60, 0.76, fill=BLUE_L)
    process_box(slide, "ExtraTrees概率\nP_tree", 1.18, 3.35, 1.60, 0.76, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 2.90, 2.58, 4.28, 3.05)
    arrow(slide, 2.90, 3.73, 4.28, 3.35)
    process_box(slide, "加权融合\nP_final = 0.55P_dnn + 0.45P_tree", 4.38, 2.90, 3.05, 0.90, fill=TEAL_L)
    arrow(slide, 7.55, 3.35, 8.20, 3.35)
    process_box(slide, "选择概率最大的位置\nargmax(P_final)", 8.30, 2.90, 2.20, 0.90, fill=BLUE_L)
    panel(slide, 1.15, 4.78, 4.65, 1.08, fill=PANEL, line=LINE)
    text_box(slide, "为什么不用树模型替代DNN", 1.35, 4.98, 2.90, 0.24, size=12.2, bold=True, color=NAVY)
    bullets(slide, ["ExtraTrees单独Macro F1约0.6867，低于DNN。", "它的价值是提供不同视角的概率补充。"], 1.35, 5.30, 4.15, 0.42, size=10.8)
    panel(slide, 6.35, 4.78, 4.65, 1.08, fill=PANEL, line=LINE)
    text_box(slide, "融合带来的效果", 6.55, 4.98, 2.30, 0.24, size=12.2, bold=True, color=NAVY)
    bullets(slide, ["去掉ExtraTrees融合后，Macro F1下降约0.0045。", "提升幅度不大，但方向稳定。"], 6.55, 5.30, 4.10, 0.42, size=10.8)
    page += 1

    slide = new_slide(prs, "10 训练流程", "训练、验证与最终评估", "先在训练集内部选择配置，再使用统一测试集报告最终效果", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "training_curve.png", 0.85, 1.45, 4.95, 3.20, title="训练曲线")
    metric_card(slide, 6.35, 1.72, 2.25, 0.88, "训练框架", "PyTorch", "CUDA GPU", accent=BLUE)
    metric_card(slide, 8.95, 1.72, 2.25, 0.88, "Batch Size", "512", "批训练", accent=TEAL)
    metric_card(slide, 6.35, 2.90, 2.25, 0.88, "优化器", "AdamW", "权重衰减", accent=ORANGE)
    metric_card(slide, 8.95, 2.90, 2.25, 0.88, "早停策略", "patience=22", "防止过拟合", accent=RED)
    bullets(slide, [
        "训练集内部再划分验证集，用于选择网络宽度、dropout和学习率。",
        "最终在测试集上只评估一次，输出混淆矩阵、分类报告和预测文件。",
        "训练过程使用label smoothing，使模型概率不过度自信。",
    ], 1.00, 5.25, 10.80, 0.98, size=12.6)
    page += 1

    slide = new_slide(prs, "11 结果分析", "最终模型的分类效果", "结果图用于说明整体准确率和各类别表现，而不是只看一个总分", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "confusion_matrix.png", 0.80, 1.45, 5.05, 3.80, title="深度学习模型混淆矩阵")
    table(slide, class_rows, ["位置", "Precision", "Recall", "F1"], 6.55, 1.58, 4.52, 2.48, widths=[0.72, 1.20, 1.20, 1.20], font_size=10.0, header_fill=TEAL)
    metric_card(slide, 6.65, 4.55, 1.55, 0.78, "PG F1", "0.861", "最高", accent=TEAL)
    metric_card(slide, 8.43, 4.55, 1.55, 0.78, "C F1", "0.755", "较稳定", accent=BLUE)
    metric_card(slide, 10.21, 4.55, 1.55, 0.78, "Macro F1", "0.7258", "整体均衡", accent=ORANGE)
    bullets(slide, [
        "PG识别最好，说明组织、三分、失误等特征组合较清晰。",
        "C/PF、PF/SF、SF/SG仍有混淆，符合篮球位置连续性的特点。",
    ], 0.95, 5.95, 10.90, 0.62, size=12.3)
    page += 1

    slide = new_slide(prs, "12 消融实验", "模块贡献：移除后性能下降", "消融实验用于回答“提升到底来自哪里”", page, total)
    image_box(slide, ABLATION_CN, 0.85, 1.42, 6.95, 4.20, title="中文标签消融实验图")
    panel(slide, 8.25, 1.60, 3.35, 3.95, fill=PANEL, line=LINE)
    section_label(slide, "读图方式", 8.55, 1.95, 1.08, fill=TEAL_L, color=TEAL)
    bullets(slide, [
        "横轴越长，说明该模块越重要。",
        "DNN主干和DNN集成贡献最大。",
        "工程特征和年代专家提供明显补充。",
        "ExtraTrees和对比学习带来小幅稳定提升。",
    ], 8.55, 2.45, 2.65, 1.10, size=11.3)
    text_box(slide, "结论：性能不是来自单个技巧，而是来自“特征 + DNN表示 + 集成 + 年代差异 + 概率融合”的组合。", 8.55, 4.50, 2.65, 0.55, size=11.3, bold=True, color=NAVY)
    page += 1

    slide = new_slide(prs, "13 汇报结论", "最终结论与讲解顺序", "答辩时按“任务—基线—模型—结果—消融”的顺序讲，逻辑最清楚", page, total)
    metric_card(slide, 0.85, 1.55, 2.15, 0.90, "Accuracy", percent(dnn["accuracy"]), "最终模型", accent=TEAL)
    metric_card(slide, 3.28, 1.55, 2.15, 0.90, "Macro F1", percent(dnn["macro_f1"]), "最终模型", accent=BLUE)
    metric_card(slide, 5.71, 1.55, 2.15, 0.90, "特征数", "45", "正式版本", accent=ORANGE)
    metric_card(slide, 8.14, 1.55, 2.15, 0.90, "合规性", "满足", "无身份泄露", accent=RED)
    table(slide, [
        ["1", "任务", "用赛季技术统计预测五类位置"],
        ["2", "基线", "前四个基础模型提供对照"],
        ["3", "模型", "45维特征 + Residual DNN + 年代专家 + 融合"],
        ["4", "结果", "DNN在Accuracy和Macro F1上最高"],
        ["5", "消融", "证明各模块都有可解释贡献"],
    ], ["顺序", "部分", "讲解要点"], 1.00, 3.05, 10.55, 2.35, widths=[0.75, 1.20, 8.60], font_size=10.4, header_fill=NAVY)
    text_box(slide, "一句话总结：深度学习模型通过更合适的特征表达和概率融合，能够更好处理NBA位置之间的连续边界。", 1.05, 6.08, 10.80, 0.35, size=13, bold=True, color=NAVY)

    prs.save(OUT_PATH)
    print(f"Saved design PPT: {OUT_PATH}")
    print(f"Saved Chinese comparison chart: {COMPARISON_CN}")
    print(f"Saved Chinese ablation chart: {ABLATION_CN}")


if __name__ == "__main__":
    main()
