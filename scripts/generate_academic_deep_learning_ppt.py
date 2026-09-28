from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUT_PATH = RESULTS / "NBA_position_deep_learning_academic_report.pptx"

WIDE = 13.333
HIGH = 7.5

FONT = "Microsoft YaHei"
FONT_EN = "Aptos"

NAVY = RGBColor(24, 43, 72)
BLUE = RGBColor(36, 99, 171)
TEAL = RGBColor(26, 144, 127)
ORANGE = RGBColor(230, 126, 34)
RED = RGBColor(198, 57, 70)
INK = RGBColor(31, 41, 55)
GREY = RGBColor(91, 103, 118)
MID = RGBColor(148, 163, 184)
LINE = RGBColor(218, 226, 237)
BG = RGBColor(247, 249, 252)
PANEL = RGBColor(255, 255, 255)
BLUE_L = RGBColor(230, 240, 252)
TEAL_L = RGBColor(226, 246, 242)
ORANGE_L = RGBColor(255, 241, 224)
RED_L = RGBColor(253, 232, 235)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def percent(value: float) -> str:
    return f"{value * 100:.2f}%"


def num(value: float) -> str:
    return f"{value:.4f}"


def fill_shape(shape, fill: RGBColor, line: RGBColor | None = None, width: float = 0.75) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if hasattr(shape, "line"):
        shape.line.color.rgb = line if line is not None else fill
        shape.line.width = Pt(width)


def set_para(paragraph, size=14, bold=False, color=INK, align=None) -> None:
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align


def text_box(slide, text, x, y, w, h, size=14, bold=False, color=INK, align=None, valign=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if valign is not None:
        tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    set_para(p, size=size, bold=bold, color=color, align=align)
    return box


def bullets(slide, items, x, y, w, h, size=13, color=INK, gap=0.93):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.line_spacing = gap
    return box


def new_slide(prs: Presentation, section: str, title: str, subtitle: str, page: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(WIDE), Inches(0.16))
    fill_shape(top, NAVY)
    text_box(slide, section, 0.55, 0.32, 2.4, 0.24, size=9, bold=True, color=TEAL)
    text_box(slide, title, 0.55, 0.56, 9.6, 0.48, size=23, bold=True, color=NAVY)
    if subtitle:
        text_box(slide, subtitle, 0.58, 1.02, 11.6, 0.28, size=10.5, color=GREY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.27), Inches(1.05), Inches(0.04))
    fill_shape(accent, TEAL)
    text_box(slide, f"{page:02d}/{total:02d}", 11.85, 7.08, 0.9, 0.22, size=9, color=GREY, align=PP_ALIGN.RIGHT)
    text_box(slide, "NBA Position Classification · Deep Learning Report", 0.55, 7.08, 5.2, 0.22, size=8.5, color=MID)
    return slide


def panel(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_shape(shape, fill, line)
    return shape


def metric_card(slide, x, y, w, h, label, value, note="", accent=BLUE, fill=PANEL):
    panel(slide, x, y, w, h, fill=fill, line=LINE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    fill_shape(bar, accent)
    text_box(slide, label, x + 0.20, y + 0.13, w - 0.35, 0.25, size=10, bold=True, color=GREY)
    text_box(slide, value, x + 0.20, y + 0.42, w - 0.30, 0.38, size=21, bold=True, color=NAVY)
    if note:
        text_box(slide, note, x + 0.20, y + 0.84, w - 0.30, 0.25, size=8.5, color=GREY)


def table(slide, rows, cols, x, y, w, h, widths=None, font_size=9.5, header_fill=NAVY):
    ppt_table = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for idx, col_width in enumerate(widths):
            ppt_table.columns[idx].width = Inches(col_width)

    for j, col in enumerate(cols):
        cell = ppt_table.cell(0, j)
        cell.text = col
        fill_shape(cell, header_fill)
        cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            set_para(p, size=font_size, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)

    for i, row in enumerate(rows):
        row_fill = RGBColor(255, 255, 255) if i % 2 == 0 else RGBColor(241, 245, 249)
        for j, value in enumerate(row):
            cell = ppt_table.cell(i + 1, j)
            cell.text = str(value)
            fill_shape(cell, row_fill, LINE)
            cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                set_para(p, size=font_size, color=INK, align=PP_ALIGN.CENTER)
    return ppt_table


def image_box(slide, path: Path, x, y, w, h, title: str | None = None, caption: str | None = None):
    panel(slide, x, y, w, h, fill=PANEL, line=LINE)
    top_pad = 0.32 if title else 0.08
    bottom_pad = 0.22 if caption else 0.08
    if title:
        text_box(slide, title, x + 0.16, y + 0.11, w - 0.32, 0.22, size=10.5, bold=True, color=NAVY)
    if caption:
        text_box(slide, caption, x + 0.16, y + h - 0.28, w - 0.32, 0.18, size=7.8, color=GREY, align=PP_ALIGN.CENTER)

    img_x = x + 0.12
    img_y = y + top_pad
    img_w = w - 0.24
    img_h = h - top_pad - bottom_pad
    if not path.exists():
        text_box(slide, f"缺失图片：{path.name}", img_x, img_y + img_h / 2 - 0.15, img_w, 0.3, size=11, color=RED, align=PP_ALIGN.CENTER)
        return None

    with Image.open(path) as img:
        iw, ih = img.size
    image_ratio = iw / ih
    box_ratio = img_w / img_h
    if image_ratio >= box_ratio:
        final_w = img_w
        final_h = img_w / image_ratio
    else:
        final_h = img_h
        final_w = img_h * image_ratio
    px = img_x + (img_w - final_w) / 2
    py = img_y + (img_h - final_h) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(final_w), height=Inches(final_h))


def process_box(slide, text, x, y, w, h, fill=BLUE_L, line=LINE, color=NAVY, size=11):
    shape = panel(slide, x, y, w, h, fill=fill, line=line)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    set_para(p, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)
    return shape


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.3):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def section_label(slide, text, x, y, w, fill=TEAL_L, color=TEAL):
    panel(slide, x, y, w, 0.30, fill=fill, line=fill)
    text_box(slide, text, x + 0.06, y + 0.06, w - 0.12, 0.18, size=8.5, bold=True, color=color, align=PP_ALIGN.CENTER)


def main() -> None:
    bayes = read_json(RESULTS / "experiment1_bayes" / "metrics.json")
    kmeans = read_json(RESULTS / "experiment2_kmeans" / "metrics.json")
    id3 = read_json(RESULTS / "experiment3_id3" / "metrics.json")
    c45 = read_json(RESULTS / "experiment4_c45" / "metrics.json")
    dnn = read_json(RESULTS / "experiment5_dnn" / "metrics.json")
    class_metrics = pd.read_csv(RESULTS / "experiment5_dnn" / "class_metrics.csv")
    fs = pd.read_csv(RESULTS / "experiment5_feature_selection_search" / "compact_feature_search_results.csv")

    class_rows = []
    for pos in ["C", "PF", "PG", "SF", "SG"]:
        sub = class_metrics[class_metrics["Position"] == pos].set_index("Metric")["Score"]
        class_rows.append([pos, num(float(sub["precision"])), num(float(sub["recall"])), num(float(sub["f1-score"]))])

    feature_rows = []
    for name in ["separation_core_45", "separation_plus_all_profiles_54", "domain_compact_42", "importance_top_35", "profile_core_32"]:
        row = fs[fs["feature_set"] == name].iloc[0]
        label = {
            "separation_core_45": "45维正式版",
            "separation_plus_all_profiles_54": "54维增强版",
            "domain_compact_42": "42维紧凑版",
            "importance_top_35": "35维重要性版",
            "profile_core_32": "32维画像版",
        }[name]
        feature_rows.append([label, int(row["feature_count"]), num(row["accuracy"]), num(row["macro_f1"])])

    prs = Presentation()
    prs.slide_width = Inches(WIDE)
    prs.slide_height = Inches(HIGH)
    total = 12
    page = 1

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), Inches(HIGH))
    fill_shape(band, NAVY)
    text_box(slide, "NBA球员位置预测实验汇报", 0.75, 0.55, 8.3, 0.58, size=30, bold=True, color=NAVY)
    text_box(slide, "重点：深度学习模型设计、训练流程与五个实验横向对比", 0.78, 1.20, 9.8, 0.36, size=15, color=GREY)
    text_box(slide, "Deep Learning Model for NBA Position Classification", 0.78, 1.62, 8.2, 0.25, size=10, color=MID)
    metric_card(slide, 9.65, 0.58, 2.65, 0.88, "Accuracy", percent(dnn["accuracy"]), "测试集", accent=TEAL)
    metric_card(slide, 9.65, 1.62, 2.65, 0.88, "Macro F1", percent(dnn["macro_f1"]), "五类平均", accent=BLUE)
    metric_card(slide, 9.65, 2.66, 2.65, 0.88, "Feature Count", str(dnn["feature_count"]), "45维正式版", accent=ORANGE)
    image_box(slide, RESULTS / "supervised_model_comparison.png", 0.78, 2.45, 5.85, 3.78, title="监督分类模型整体对比")
    image_box(slide, RESULTS / "experiment5_dnn" / "confusion_matrix.png", 6.86, 3.12, 5.45, 3.11, title="最终模型混淆矩阵")
    text_box(slide, "结论先行：深度学习融合模型在不使用 Player、Tm 或历史位置等身份信息的前提下，显著超过前四个基础模型。", 0.82, 6.58, 11.2, 0.34, size=12.5, bold=True, color=NAVY)
    text_box(slide, f"{page:02d}/{total:02d}", 11.85, 7.08, 0.9, 0.22, size=9, color=GREY, align=PP_ALIGN.RIGHT)
    page += 1

    # 2. Research task
    slide = new_slide(prs, "01 问题定义", "研究任务与实验流程", "从原始 NBA 赛季技术统计出发，构建五类位置预测任务", page, total)
    process_box(slide, "原始数据\nNBA_Season_Stats.csv", 0.75, 1.75, 2.05, 0.72, fill=BLUE_L)
    arrow(slide, 2.88, 2.10, 3.25, 2.10)
    process_box(slide, "合规清洗\n去除 Player / Tm", 3.32, 1.75, 2.05, 0.72, fill=TEAL_L)
    arrow(slide, 5.45, 2.10, 5.82, 2.10)
    process_box(slide, "特征构建\n原始26维 / 深度45维", 5.90, 1.75, 2.35, 0.72, fill=BLUE_L)
    arrow(slide, 8.34, 2.10, 8.71, 2.10)
    process_box(slide, "模型训练\n五个实验", 8.78, 1.75, 1.80, 0.72, fill=TEAL_L)
    arrow(slide, 10.66, 2.10, 11.03, 2.10)
    process_box(slide, "统一评估\nAccuracy / F1", 11.10, 1.75, 1.55, 0.72, fill=BLUE_L)

    metric_card(slide, 0.78, 3.08, 2.65, 1.05, "分类目标", "5类", "C / PF / PG / SF / SG", accent=BLUE)
    metric_card(slide, 3.72, 3.08, 2.65, 1.05, "训练集", "14,981", "传统实验外层训练", accent=TEAL)
    metric_card(slide, 6.66, 3.08, 2.65, 1.05, "测试集", "3,746", "统一测试评估", accent=BLUE)
    metric_card(slide, 9.60, 3.08, 2.65, 1.05, "DNN验证集", "2,997", "训练集内部划分", accent=TEAL)
    bullets(slide, [
        "前四个实验保持基础算法版本，用作可解释的基线对照。",
        "实验五在同一任务上引入深度模型、特征工程、年代专家和概率融合。",
        "合规约束：不使用球员姓名、球队、历史位置等身份或泄露信息。",
    ], 0.85, 4.80, 11.7, 1.15, size=13.5)
    page += 1

    # 3. Baselines
    slide = new_slide(prs, "02 基线实验", "前四个模型：方法定位与局限", "基础模型不是为了堆叠调参，而是提供清晰对照：统计假设、距离聚类、单树划分", page, total)
    rows = [
        ["贝叶斯", "GaussianNB", "条件独立 + 高斯分布", f"Acc {num(bayes['accuracy'])}\nMacro F1 {num(bayes['macro_f1'])}"],
        ["K-Means", "无监督聚类", "欧氏距离形成5个簇", f"Silhouette {num(kmeans['silhouette_score'])}\nARI {num(kmeans['adjusted_rand_index'])}"],
        ["ID3", "信息增益", "离散特征上的单树规则", f"Acc {num(id3['accuracy'])}\nMacro F1 {num(id3['macro_f1'])}"],
        ["C4.5", "信息增益率", "修正多取值偏好", f"Acc {num(c45['accuracy'])}\nMacro F1 {num(c45['macro_f1'])}"],
    ]
    table(slide, rows, ["实验", "算法", "核心假设 / 划分依据", "测试结果"], 0.65, 1.55, 6.05, 3.0, widths=[0.95, 1.35, 2.25, 1.5], font_size=9.2)
    image_box(slide, RESULTS / "experiment2_kmeans" / "pca_clusters.png", 7.05, 1.45, 2.65, 2.20, title="K-Means PCA簇")
    image_box(slide, RESULTS / "experiment2_kmeans" / "cluster_position_heatmap.png", 9.95, 1.45, 2.65, 2.20, title="簇-位置交叉分布")
    section_label(slide, "基线结论", 7.10, 4.08, 1.10, fill=ORANGE_L, color=ORANGE)
    bullets(slide, [
        "贝叶斯受特征相关性影响，得分、出手、时间等统计并非独立。",
        "K-Means不使用标签，适合探索结构，但很难直接恢复真实位置。",
        "ID3 / C4.5为单树模型，规则可解释，但对连续边界和摇摆位置表达不足。",
    ], 7.10, 4.50, 5.35, 1.35, size=12.5)
    page += 1

    # 4. Overall comparison
    slide = new_slide(prs, "02 基线实验", "五个实验整体指标对比", "实验五的提升不仅体现在 Accuracy，也体现在更均衡的 Macro F1", page, total)
    image_box(slide, RESULTS / "supervised_model_comparison.png", 0.70, 1.40, 6.40, 4.15, title="监督分类模型 Accuracy / Macro F1 / Weighted F1")
    rows = [
        ["贝叶斯", num(bayes["accuracy"]), num(bayes["macro_f1"]), "26"],
        ["ID3", num(id3["accuracy"]), num(id3["macro_f1"]), "26"],
        ["C4.5", num(c45["accuracy"]), num(c45["macro_f1"]), "26"],
        ["DNN融合", num(dnn["accuracy"]), num(dnn["macro_f1"]), "45"],
    ]
    table(slide, rows, ["模型", "Accuracy", "Macro F1", "特征数"], 7.42, 1.55, 4.55, 1.95, widths=[1.30, 1.10, 1.15, 0.80], font_size=10.5, header_fill=BLUE)
    metric_card(slide, 7.45, 4.10, 1.90, 0.88, "vs 贝叶斯", f"+{dnn['macro_f1'] - bayes['macro_f1']:.3f}", "Macro F1", accent=TEAL)
    metric_card(slide, 9.55, 4.10, 1.90, 0.88, "vs ID3", f"+{dnn['macro_f1'] - id3['macro_f1']:.3f}", "Macro F1", accent=TEAL)
    metric_card(slide, 11.65, 4.10, 1.15, 0.88, "领先", "显著", "同任务", accent=ORANGE)
    bullets(slide, [
        "深度学习模型对复杂非线性边界的拟合能力更强。",
        "Macro F1接近Accuracy，说明提升不是依赖某一单独类别。",
    ], 7.45, 5.38, 4.95, 0.80, size=12.2)
    page += 1

    # 5. DNN framework
    slide = new_slide(prs, "03 深度学习模型", "实验五总体框架", "45维技术特征进入多分支模型，最终用概率融合得到五类位置预测", page, total)
    process_box(slide, "45维核心技术特征", 0.70, 2.70, 1.95, 0.70, fill=TEAL_L)
    arrow(slide, 2.72, 3.05, 3.18, 3.05)
    process_box(slide, "标准化\nStandardScaler", 3.25, 2.70, 1.55, 0.70, fill=BLUE_L)
    arrow(slide, 4.88, 3.05, 5.30, 2.35)
    arrow(slide, 4.88, 3.05, 5.30, 3.05)
    arrow(slide, 4.88, 3.05, 5.30, 3.78)
    process_box(slide, "全局DNN集成\n3 Residual + 3 Contrastive", 5.40, 1.83, 2.35, 0.78, fill=BLUE_L)
    process_box(slide, "年代专家DNN\n1980s / 1990s / 2000s / 2010s+", 5.40, 2.82, 2.35, 0.78, fill=TEAL_L)
    process_box(slide, "ExtraTrees\n辅助概率模型", 5.40, 3.81, 2.35, 0.78, fill=ORANGE_L, color=ORANGE)
    arrow(slide, 7.85, 2.22, 8.52, 2.80)
    arrow(slide, 7.85, 3.20, 8.52, 3.20)
    arrow(slide, 7.85, 4.20, 8.52, 3.58)
    process_box(slide, "概率融合\nDNN + Tree", 8.62, 2.80, 1.78, 0.78, fill=TEAL_L)
    arrow(slide, 10.48, 3.20, 10.95, 3.20)
    process_box(slide, "最终预测\nC / PF / PG / SF / SG", 11.05, 2.80, 1.62, 0.78, fill=BLUE_L)
    bullets(slide, [
        "主干模型：Residual DNN学习非线性统计组合。",
        "监督对比学习：辅助拉开相邻位置的隐空间距离。",
        "年代专家：吸收不同时期位置职责和三分使用方式变化。",
        "ExtraTrees：提供与DNN不同错误模式的概率补充。",
    ], 0.88, 5.18, 11.4, 1.05, size=12.8)
    page += 1

    # 6. Feature system
    slide = new_slide(prs, "03 深度学习模型", "45维核心特征体系与筛选结果", "目标不是无限增加字段，而是在更少输入下保留关键位置判别信息", page, total)
    families = [
        ("效率", "FG% / 3P% / eFG% / FT%", BLUE_L, BLUE),
        ("使用强度", "PPG / MPG / PTS_36 / FTA_36", TEAL_L, TEAL),
        ("组织防守", "AST_TOV / STL_36 / BLK_36", BLUE_L, BLUE),
        ("篮板画像", "ORB_Ratio / DRB_Ratio / TRB_36", TEAL_L, TEAL),
        ("位置分离", "Center_PF / SF_PF / SG_PG", ORANGE_L, ORANGE),
    ]
    for i, (name, desc, fill, color) in enumerate(families):
        y = 1.55 + i * 0.72
        panel(slide, 0.72, y, 4.35, 0.52, fill=fill, line=fill)
        text_box(slide, name, 0.90, y + 0.12, 0.80, 0.18, size=9.2, bold=True, color=color, align=PP_ALIGN.CENTER)
        text_box(slide, desc, 1.75, y + 0.11, 3.05, 0.20, size=9.6, color=INK)
    metric_card(slide, 0.75, 5.45, 1.55, 0.78, "候选特征", "76", "原始+工程", accent=BLUE)
    metric_card(slide, 2.48, 5.45, 1.55, 0.78, "正式输入", "45", "减少40.8%", accent=TEAL)
    metric_card(slide, 4.20, 5.45, 1.55, 0.78, "Macro F1", percent(dnn["macro_f1"]), "基本保持", accent=ORANGE)
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "feature_count_vs_macro_f1.png", 6.05, 1.42, 3.15, 2.40, title="特征数量与Macro F1")
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "compact_feature_search_macro_f1.png", 9.43, 1.42, 3.15, 2.40, title="紧凑特征集比较")
    table(slide, feature_rows, ["特征集", "维度", "Acc", "Macro F1"], 6.10, 4.28, 6.18, 1.75, widths=[2.25, 0.80, 1.05, 1.20], font_size=9.2, header_fill=TEAL)
    page += 1

    # 7. Residual DNN
    slide = new_slide(prs, "03 深度学习模型", "Residual DNN 与监督对比学习模块", "用残差连接稳定训练，用对比目标增强相邻位置区分能力", page, total)
    process_box(slide, "Input\n45维", 0.85, 2.05, 1.08, 0.70, fill=BLUE_L)
    arrow(slide, 2.00, 2.40, 2.42, 2.40)
    process_box(slide, "Linear + BN\nReLU + Dropout", 2.50, 1.93, 1.55, 0.95, fill=TEAL_L)
    arrow(slide, 4.12, 2.40, 4.52, 2.40)
    process_box(slide, "Residual Block ×2\nx + F(x)", 4.60, 1.93, 1.65, 0.95, fill=BLUE_L)
    arrow(slide, 6.32, 2.40, 6.72, 2.40)
    process_box(slide, "Classifier\n5类概率", 6.80, 1.93, 1.35, 0.95, fill=TEAL_L)
    arrow(slide, 8.20, 2.40, 8.60, 2.40)
    process_box(slide, "Prediction\nC/PF/PG/SF/SG", 8.68, 2.05, 1.62, 0.70, fill=BLUE_L)
    arrow(slide, 6.98, 2.93, 6.98, 3.38)
    process_box(slide, "Projection Head\nSupervised Contrastive", 5.95, 3.48, 2.10, 0.78, fill=ORANGE_L, color=ORANGE)
    text_box(slide, "Loss = CrossEntropy(label_smoothing=0.02) + 0.1 × SupervisedContrastiveLoss", 0.95, 4.75, 10.4, 0.32, size=14.5, bold=True, color=NAVY)
    rows = [
        ["基础成员", "Residual DNN × 3", "不同随机种子降低偶然性"],
        ["增强成员", "Contrastive Residual DNN × 3", "隐空间拉开类别边界"],
        ["优化", "AdamW + CosineAnnealingLR", "训练稳定、权重衰减抑制过拟合"],
        ["批大小", str(dnn["batch_size"]), "适合GPU批训练"],
    ]
    table(slide, rows, ["模块", "设置", "作用"], 0.95, 5.35, 11.20, 1.16, widths=[1.55, 3.15, 6.50], font_size=9.2, header_fill=BLUE)
    page += 1

    # 8. Training and fusion
    slide = new_slide(prs, "03 深度学习模型", "训练、验证与概率融合流程", "先在验证集选择稳定配置，再全训练集重训并在测试集统一评估", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "training_curve.png", 0.70, 1.42, 4.25, 3.10, title="训练/验证曲线")
    image_box(slide, RESULTS / "experiment5_dnn" / "validation_comparison.png", 5.20, 1.42, 4.15, 3.10, title="验证配置对比")
    metric_card(slide, 9.70, 1.55, 2.55, 0.82, "GPU", "RTX 4060", "PyTorch cuda", accent=TEAL)
    metric_card(slide, 9.70, 2.55, 2.55, 0.82, "训练策略", "Early Stop", f"patience=22", accent=BLUE)
    metric_card(slide, 9.70, 3.55, 2.55, 0.82, "融合方式", "Soft Prob.", "概率平均", accent=ORANGE)
    bullets(slide, [
        "全局DNN概率与年代专家概率先融合，再与ExtraTrees概率加权融合。",
        "采用软概率而非硬投票，使模型保留类别不确定性。",
        "最终模型在测试集上统一输出 predictions、classification report 和可视化图。",
    ], 0.88, 5.08, 11.35, 1.15, size=12.8)
    page += 1

    # 9. Classification effect
    slide = new_slide(prs, "04 结果分析", "实验五测试集分类效果", "PG和C最稳定，PF/SF/SG等摇摆位置是主要混淆区域", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "class_metrics.png", 0.65, 1.38, 4.00, 3.45, title="各类别 Precision / Recall / F1")
    image_box(slide, RESULTS / "experiment5_dnn" / "confusion_matrix.png", 4.95, 1.38, 4.05, 3.45, title="混淆矩阵")
    table(slide, class_rows, ["位置", "Precision", "Recall", "F1"], 9.35, 1.50, 3.00, 2.25, widths=[0.60, 0.80, 0.80, 0.80], font_size=9.4, header_fill=TEAL)
    metric_card(slide, 9.42, 4.15, 1.35, 0.78, "PG F1", "0.861", "最高", accent=TEAL)
    metric_card(slide, 10.95, 4.15, 1.35, 0.78, "C F1", "0.755", "稳定", accent=BLUE)
    bullets(slide, [
        "PG类别识别最好：助攻、三分、失误、出场方式等特征组合较清晰。",
        "C/PF、PF/SF、SF/SG之间仍有混淆，符合篮球位置连续性的现实特征。",
    ], 0.85, 5.35, 11.35, 0.80, size=12.5)
    page += 1

    # 10. Error interpretation
    slide = new_slide(prs, "04 结果分析", "错误模式解释：为什么不是所有类别都同样容易", "篮球位置并非严格离散标签，前锋和摇摆人边界天然更模糊", page, total)
    image_box(slide, RESULTS / "experiment5_dnn" / "prediction_distribution.png", 0.72, 1.42, 3.45, 2.35, title="预测类别分布")
    image_box(slide, RESULTS / "experiment1_bayes" / "confusion_matrix.png", 4.45, 1.42, 3.45, 2.35, title="贝叶斯混淆矩阵")
    image_box(slide, RESULTS / "experiment5_dnn" / "confusion_matrix.png", 8.18, 1.42, 3.45, 2.35, title="DNN混淆矩阵")
    section_label(slide, "关键混淆", 0.85, 4.30, 1.05, fill=RED_L, color=RED)
    bullets(slide, [
        "C → PF：162，PF → C：148，内线职责存在重叠。",
        "SF → SG：118，SG → SF：118，锋卫摇摆位置在统计上接近。",
        "PG → SG：88，主要来自得分型后卫与组织型后卫的边界。",
    ], 0.85, 4.78, 4.95, 1.10, size=12.2)
    section_label(slide, "模型优势", 6.20, 4.30, 1.05, fill=TEAL_L, color=TEAL)
    bullets(slide, [
        "DNN没有消除所有模糊边界，但显著减少了基础模型的系统性偏差。",
        "残差结构、工程特征和概率融合共同提升了相邻位置判别能力。",
        "错误主要集中于篮球语义上相邻的类别，说明模型输出具有可解释性。",
    ], 6.20, 4.78, 5.95, 1.10, size=12.2)
    page += 1

    # 11. Ablation
    slide = new_slide(prs, "04 结果分析", "消融实验：性能提升来自哪些模块", "逐项移除模块后比较 Macro F1 下降，验证方案不是单点技巧", page, total)
    image_box(slide, RESULTS / "experiment5_ablation" / "ablation_delta_macro_f1.png", 0.70, 1.40, 5.20, 3.55, title="移除模块后的Macro F1下降")
    image_box(slide, RESULTS / "experiment5_ablation" / "ablation_macro_f1.png", 6.18, 1.40, 3.15, 3.55, title="各消融版本Macro F1")
    rows = [
        ["只用ExtraTrees", "-0.0393", "DNN主干贡献最大"],
        ["单个Residual DNN", "-0.0340", "集成降低随机性"],
        ["去掉工程特征", "-0.0235", "篮球语义特征有效"],
        ["去掉年代专家", "-0.0176", "时代差异重要"],
        ["去掉ExtraTrees", "-0.0045", "辅助概率补充"],
    ]
    table(slide, rows, ["消融项", "Macro F1变化", "解释"], 9.60, 1.55, 2.65, 3.22, widths=[1.05, 0.75, 0.85], font_size=8.2, header_fill=ORANGE)
    bullets(slide, [
        "最核心贡献：DNN集成、工程特征、年代专家。",
        "小幅但稳定贡献：ExtraTrees融合、Contrastive成员、监督对比损失。",
    ], 0.90, 5.35, 11.10, 0.78, size=12.6)
    page += 1

    # 12. Conclusion
    slide = new_slide(prs, "05 结论", "结论与答辩要点", "实验五在合规约束下取得最优结果，并用对比和消融说明提升来源", page, total)
    metric_card(slide, 0.78, 1.55, 2.20, 0.95, "最终Accuracy", percent(dnn["accuracy"]), "测试集", accent=TEAL)
    metric_card(slide, 3.25, 1.55, 2.20, 0.95, "最终Macro F1", percent(dnn["macro_f1"]), "五类平均", accent=BLUE)
    metric_card(slide, 5.72, 1.55, 2.20, 0.95, "正式特征", "45维", "比76维少40.8%", accent=ORANGE)
    metric_card(slide, 8.19, 1.55, 2.20, 0.95, "任务合规", "无身份字段", "不使用Player/Tm", accent=RED)
    image_box(slide, RESULTS / "supervised_model_comparison.png", 0.82, 3.05, 4.85, 2.70, title="最终横向对比")
    image_box(slide, RESULTS / "experiment5_feature_selection_search" / "feature_count_vs_macro_f1.png", 6.00, 3.05, 3.05, 2.70, title="特征筛选证据")
    section_label(slide, "汇报主线", 9.55, 3.03, 1.10, fill=TEAL_L, color=TEAL)
    bullets(slide, [
        "先说明任务和合规数据处理，避免老师质疑数据泄露。",
        "再说明前四个基础模型作为对照，指标低是算法假设限制。",
        "重点讲实验五：45维特征、Residual DNN、年代专家、概率融合。",
        "最后用总体对比、混淆矩阵和消融实验证明优势来源。",
    ], 9.55, 3.48, 3.05, 1.70, size=11.4)
    text_box(slide, "核心结论：深度学习模型不是单纯堆参数，而是把篮球语义特征、非线性表示、时代差异和概率融合组合起来，因此能稳定超过传统基础模型。", 0.90, 6.25, 11.25, 0.38, size=13, bold=True, color=NAVY)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Saved academic PPT: {OUT_PATH}")


if __name__ == "__main__":
    main()
